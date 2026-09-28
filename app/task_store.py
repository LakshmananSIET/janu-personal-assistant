from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook, load_workbook


DATA_DIR = Path("data")
TASK_FILE = DATA_DIR / "janu_trial_tasks.xlsx"

HEADERS = [
    "id",
    "created_at",
    "task",
    "due_date",
    "deadline",
    "status",
    "source_text",
]


def save_task(task: dict) -> int:
    DATA_DIR.mkdir(exist_ok=True)

    if TASK_FILE.exists():
        workbook = load_workbook(TASK_FILE)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Tasks"
        sheet.append(HEADERS)

    next_id = sheet.max_row
    sheet.append([
        next_id,
        __import__("datetime").datetime.now().isoformat(timespec="seconds"),
        task.get("task"),
        task.get("due_date"),
        task.get("deadline"),
        task.get("status", "pending"),
        task.get("source_text"),
    ])

    workbook.save(TASK_FILE)
    return next_id
